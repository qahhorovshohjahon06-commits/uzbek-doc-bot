import io
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN


def build_docx(title: str, content: str, doc_type: str) -> io.BytesIO:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.runs[0]
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    doc.add_paragraph()

    lines = content.strip().split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue

        if re.match(r"^\d+\.\s+[A-ZA-Z\u0400-\u04FF]", line) or (
            line.isupper() and len(line) > 5
        ):
            heading = doc.add_heading(line, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = heading.runs[0]
            run.font.size = Pt(13)
            run.font.bold = True
        elif line.startswith("- ") or line.startswith("• "):
            para = doc.add_paragraph(line[2:], style="List Bullet")
            para.paragraph_format.left_indent = Inches(0.3)
        else:
            para = doc.add_paragraph(line)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.first_line_indent = Inches(0.5)
            para.paragraph_format.space_after = Pt(6)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(0.8)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_pptx(title: str, content: str) -> io.BytesIO:
    prs = Presentation()
    prs.slide_width = PInches(13.33)
    prs.slide_height = PInches(7.5)

    DARK_BG = PRGBColor(0x1A, 0x1A, 0x2E)
    ACCENT = PRGBColor(0x16, 0x21, 0x3E)
    WHITE = PRGBColor(0xFF, 0xFF, 0xFF)
    GOLD = PRGBColor(0xE9, 0x4F, 0x37)

    slide_layout = prs.slide_layouts[6]

    def add_background(slide, color=DARK_BG):
        from pptx.util import Emu
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_title_slide(slide, main_title):
        add_background(slide)
        txBox = slide.shapes.add_textbox(
            PInches(1), PInches(2.5), PInches(11.33), PInches(1.5)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = main_title
        run.font.size = PPt(40)
        run.font.bold = True
        run.font.color.rgb = WHITE

        line_box = slide.shapes.add_textbox(
            PInches(4), PInches(4.2), PInches(5.33), PInches(0.1)
        )
        line_tf = line_box.text_frame
        lp = line_tf.paragraphs[0]
        lr = lp.add_run()
        lr.text = "─" * 40
        lr.font.color.rgb = GOLD
        lr.font.size = PPt(12)

    def add_content_slide(slide, slide_title, bullets):
        add_background(slide)

        title_box = slide.shapes.add_textbox(
            PInches(0.5), PInches(0.3), PInches(12.33), PInches(0.9)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = slide_title
        run.font.size = PPt(28)
        run.font.bold = True
        run.font.color.rgb = GOLD

        content_box = slide.shapes.add_textbox(
            PInches(0.7), PInches(1.4), PInches(11.93), PInches(5.5)
        )
        ctf = content_box.text_frame
        ctf.word_wrap = True

        for i, bullet in enumerate(bullets):
            if i == 0:
                cp = ctf.paragraphs[0]
            else:
                cp = ctf.add_paragraph()
            cp.space_before = PPt(8)
            cr = cp.add_run()
            cr.text = f"▸  {bullet}"
            cr.font.size = PPt(20)
            cr.font.color.rgb = WHITE

    slide_blocks = re.split(r"SLAYD\s+\d+:", content, flags=re.IGNORECASE)
    slide_blocks = [b.strip() for b in slide_blocks if b.strip()]

    if not slide_blocks:
        slide_blocks = [content]

    first_slide = prs.slides.add_slide(slide_layout)
    add_title_slide(first_slide, title)

    for block in slide_blocks:
        lines = [l.strip() for l in block.strip().split("\n") if l.strip()]
        if not lines:
            continue

        slide_title = lines[0]
        bullets = []
        for line in lines[1:]:
            if line.startswith("- ") or line.startswith("• "):
                bullets.append(line[2:])
            elif line:
                bullets.append(line)

        if not bullets:
            continue

        content_slide = prs.slides.add_slide(slide_layout)
        add_content_slide(content_slide, slide_title, bullets[:6])

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
